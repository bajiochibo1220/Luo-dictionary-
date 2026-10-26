"""One self-contained training presentation, with a matching PDF and PowerPoint."""
import sys, json, importlib.util, copy, re, zipfile, hashlib
from pathlib import Path
BASE = Path(__file__).parent
sys.path.insert(0, str(BASE / 'training-python'))
spec = importlib.util.spec_from_file_location('luo_training_base', BASE / 'build_luo_training.py')
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.pdfgen import canvas
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image, ImageDraw

OUT = Path('D:/my-business-website/deliverables/Luo-Heritage-One-Presentation')
OUT.mkdir(parents=True, exist_ok=True)
m.OUT = OUT
STEM = 'Luo-Heritage-Complete-Training'
COMMON = [m.LOGIN, m.RESET, m.NEW_PASSWORD, m.MOBILE, m.ACCOUNT, m.PRIVACY, m.HELP, m.FLOW,
          m.SOURCE, m.ATTACH, m.SAVE, m.STUDY, m.OFFLINE, m.ENROLL, m.CLASS_JOIN, m.CHAT]
common_titles = {s['title'] for s in COMMON}
pages = [dict(kind='cover'), dict(kind='welcome'), dict(kind='contents')]
common_start = 4
for item in COMMON:
    pages.append(dict(kind='task', section='Common steps', content=copy.deepcopy(item)))
sections = []
flows = [
    ['Find material', 'Read or listen', 'Save or contribute'],
    ['Join a course', 'Complete work', 'Read feedback'],
    ['Create a class', 'Set assignments', 'Review learning'],
    ['Prepare draft', 'Check permission', 'Submit for review'],
    ['Read source', 'Check accuracy', 'Record decision'],
    ['Read feedback', 'Correct draft', 'Return for review'],
    ['Check review', 'Check permission', 'Release correctly'],
    ['Check language', 'Coordinate staff', 'Follow up work'],
    ['Oversee languages', 'Support staff', 'Check activity'],
    ['Protect ownership', 'Assign access', 'Keep records safe'],
    ['Confirm access', 'Read sources', 'Respect conditions'],
]
for i, guide in enumerate(m.decks):
    section = dict(role=guide['title'], start=len(pages)+1, goal=guide['goal'], stem=guide['stem'])
    sections.append(section)
    divider = m.slide(guide['title'], [guide['goal'],
        f'Read the common steps on slides {common_start}–{common_start+len(COMMON)-1} when you need them. Then follow the tasks in this section.',
        'Use only the duties and language assigned to your own account. Choose Contents at the bottom of a slide to return to the role list.'
    ], diagram=flows[i], result='Read a step, try it, and check the result before moving on.')
    pages.append(dict(kind='divider', section=guide['title'], content=divider))
    # These eight opening/closing pages were repeated in every separate guide.
    # Keep shared task instructions once, while preserving every role-specific task.
    for item in guide['slides'][5:-3]:
        if item['title'] in common_titles:
            continue
        item = copy.deepcopy(item)
        if item.get('tip') == 'This guide does not promise a self-service private-dataset download portal.':
            item['tip'] = 'Your assigned role and language determine which records you may open.'
        pages.append(dict(kind='task', section=guide['title'], content=item))
    closing = copy.deepcopy(guide['slides'][-1])
    closing['title'] = 'Check your skills before working alone'
    pages.append(dict(kind='check', section=guide['title'], content=closing))
    section['end'] = len(pages)

pages.append(dict(kind='task', section='Help and next steps', content=m.slide('Finish confidently and ask for help', [
    'Keep this one file. Return to Contents whenever you need another task or role.',
    'If a screen is empty, access is refused or a service is unavailable, read the message and contact the person who manages your account.',
    'Give support the screen name, button and message. Never send your password, a reset link or private research material.'
], diagram=['Read the message', 'Check your access', 'Ask for help'], result='Use your own account and follow the conditions attached to every source.')))

total = len(pages)
p = Presentation()
p.slide_width = Pt(m.W)
p.slide_height = Pt(m.H)
p.core_properties.title = 'Luo Heritage — Complete User Training'
p.core_properties.subject = 'One presentation for all user roles; illustrated self-study training'
p.core_properties.author = 'Luo Heritage'
p.core_properties.keywords = 'user training, all roles, screenshots, self-study'
pdf = canvas.Canvas(str(OUT / (STEM+'.pdf')), pagesize=(m.W,m.H), pageCompression=1)
pdf.setTitle('Luo Heritage — Complete User Training')
pdf.setAuthor('Luo Heritage')
nav_links = []

for number, page in enumerate(pages, 1):
    d = m.SlideDraw(p,pdf)
    pdf.bookmarkPage('s'+str(number))
    if page['kind'] == 'cover':
        d.rect(0,0,960,540,m.BROWN)
        d.rect(0,0,18,540,m.GOLD)
        d.image(m.LOGO,58,42,90,90)
        d.text('LUO HERITAGE',175,62,690,21,m.GOLD,bold=True)
        d.text('LANGUAGE • CULTURE • LIVING MEMORY',175,96,690,11,'#e7d8bf')
        d.text('Complete user training',58,177,840,43,'#fffaf0',font='Georgia',leading=51)
        d.text('One presentation for every user role',60,246,820,25,'#f3e6ce',leading=31)
        d.rect(58,329,840,112,'#6d4930',radius=8)
        d.text('Start with the common steps. Then go to your role.\nRead the action, try it, and check the result.',78,347,795,21,'#fff7e8',leading=30)
        d.text('SCREENSHOTS • CLEAR STEPS • ROLE SECTIONS',60,481,780,12,m.GOLD,bold=True)
        d.text('26 October 2026  |  Combined edition',60,508,820,11,'#e7d8bf')
        d.notes('One self-contained presentation. Screenshots use labelled practice details from actual application components; they are not live-account acceptance evidence.')
        pdf.addOutlineEntry('Complete user training','s1',0)
    elif page['kind'] == 'welcome':
        content=m.slide('How to use this one presentation', [
            'This file contains all user roles. You do not need another guide, a ZIP file or a folder of screenshots to read it.',
            'Read the common steps first. Then use Contents to open the section for the work you have actually been assigned.',
            'The screenshots use practice details. Your own names, records and menus may look different. Use approved practice material when trying a task.'
        ], diagram=['Read one step','Try it yourself','Check the result'], result='Use the PDF for easy reading on a phone or computer. The PowerPoint contains the same slides.', tip='Never publish, remove accounts or change live settings just to practise.')
        page['content']=content
        m.draw_page(d,dict(title='Start here',stem=STEM),content,number,total)
    elif page['kind'] == 'contents':
        d.text('ONE PRESENTATION / ALL USER ROLES',36,22,888,10,m.MUTED,bold=True)
        d.text('Choose your section',36,47,888,29,m.BROWN,font='Georgia')
        d.text('Select a box, or go to the slide number shown.',36,92,888,16,m.INK)
        entries=[dict(role='Common steps — everyone',start=4,end=3+len(COMMON))]+sections
        for index, entry in enumerate(entries):
            column=index%2; row=index//2
            x=36+column*456; y=133+row*54
            d.rect(x,y,432,45,m.GREEN,radius=5)
            box=d.p.shapes[-1]
            label=entry['role']
            d.text(label,x+13,y+11,330,15.5,'#ffffff',bold=True)
            text=d.p.shapes[-1]
            d.text(str(entry['start']),x+367,y+10,52,16,'#ffffff',bold=True)
            nav_links.extend([(box,entry['start']),(text,entry['start']),(d.p.shapes[-1],entry['start'])])
            pdf.linkAbsolute(label,'s'+str(entry['start']),(x,540-y-45,x+432,540-y),thickness=0)
        d.text('Missing a button? Ask your administrator to check your assigned role and language.',36,468,888,13,m.MUTED)
        d.footer('Contents',number,total)
        d.notes('All buttons link within this presentation. No separate files are needed. Slide numbers also work when a viewer does not support clickable contents.')
        pdf.addOutlineEntry('Contents','s3',0)
    else:
        content=page['content']
        m.draw_page(d,dict(title=page['section'],stem=STEM),content,number,total)
        if page['kind']=='divider':
            pdf.addOutlineEntry(page['section'],'s'+str(number),0)
        elif number==4:
            pdf.addOutlineEntry('Common steps','s4',0)
            pdf.addOutlineEntry(content['title'],'s4',1)
        else:
            pdf.addOutlineEntry(content['title'],'s'+str(number),1)
    pdf.showPage()

for page in p.slides:
    for shape in page.shapes:
        if shape.name=='Contents link':
            shape.click_action.target_slide=p.slides[2]
for shape, number in nav_links:
    shape.click_action.target_slide=p.slides[number-1]
pdf.save()
p.save(str(OUT/(STEM+'.pptx')))
(OUT/'section-index.json').write_text(json.dumps(dict(totalSlides=total,common=dict(start=4,end=3+len(COMMON)),roles=sections),ensure_ascii=False,indent=2),encoding='utf8')
(OUT/'slide-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf8')
(OUT/'READ_ME.txt').write_text('LUO HERITAGE — ONE COMPLETE TRAINING PRESENTATION\n\nSEND THIS SINGLE FILE FOR EASY READING:\nLuo-Heritage-Complete-Training.pdf\n\nIt includes all user roles, screenshots and diagrams. No ZIP, extra guide or screenshot folder is required.\nFor presenting or editing, use Luo-Heritage-Complete-Training.pptx; it contains the same slides.\nOpen Contents on slide 3 and choose a role, or use the printed slide numbers.\n\n26 October 2026 — Combined edition\n',encoding='utf8')

# Verify every page and embedded asset before distribution.
pdf_read=PdfReader(OUT/(STEM+'.pdf'))
reopen=Presentation(OUT/(STEM+'.pptx'))
assert len(pdf_read.pages)==len(reopen.slides)==total
norm=lambda value:re.sub(r'\s+',' ',value).strip()
screens=0
for i, page in enumerate(pdf_read.pages):
    ppt_slide=reopen.slides[i]
    assert all(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=reopen.slide_width+10000 and shape.top+shape.height<=reopen.slide_height+10000 for shape in ppt_slide.shapes),(i+1,'shape outside slide')
    pdf_text=norm(page.extract_text())
    ppt_text=norm(' '.join(shape.text for shape in ppt_slide.shapes if shape.has_text_frame))
    expected=pages[i].get('content')
    if expected:
        for text in [expected['title']]+expected['steps']+([expected['tip']] if expected.get('tip') else []):
            assert norm(text) in pdf_text,(i+1,'PDF text missing',text)
            assert norm(text) in ppt_text,(i+1,'PowerPoint text missing',text)
        if expected.get('image'):
            screens+=1
            assert len([shape for shape in ppt_slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE])==1
    if i:
        links=[shape for shape in ppt_slide.shapes if shape.name=='Contents link']
        assert len(links)==1 and links[0].click_action.target_slide==reopen.slides[2]
with zipfile.ZipFile(OUT/(STEM+'.pptx')) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith('.rels'):
            rel=archive.read(name).decode('utf8')
            assert not ('TargetMode="External"' in rel and '/image"' in rel),'External image dependency'
            assert not re.search(r'Target="[^\"]*\.(?:pptx|pdf)"',rel),'Separate file dependency'

doc=pdfium.PdfDocument(str(OUT/(STEM+'.pdf')))
preview=OUT/'quality-previews';preview.mkdir(exist_ok=True)
thumbs=[]
for i in range(total):
    image=doc[i].render(scale=.28).to_pil().convert('RGB')
    assert image.getbbox()
    thumb=Image.new('RGB',(280,173),'white');thumb.paste(image,(5,3));ImageDraw.Draw(thumb).text((7,157),str(i+1),fill='black');thumbs.append(thumb)
    if i in [0,2,3,sections[2]['start']-1,sections[9]['start'],total-1]:
        doc[i].render(scale=1.25).to_pil().save(preview/f'slide-{i+1:03}.png')
for start in range(0,total,36):
    subset=thumbs[start:start+36];contact=Image.new('RGB',(4*280,((len(subset)+3)//4)*173),'#dedbd2')
    for index,thumb in enumerate(subset):contact.paste(thumb,((index%4)*280,(index//4)*173))
    contact.save(preview/f'contact-{start+1:03}.png')
report=dict(date='2026-10-26',slides=total,roles=len(sections),screenshotSlides=screens,allPagesRendered=True,allVisibleInstructionsVerified=True,embeddedImages=True,noSeparateFileDependencies=True,internalContentsLinks=True,files={ext:dict(bytes=(OUT/(STEM+'.'+ext)).stat().st_size,sha256=hashlib.sha256((OUT/(STEM+'.'+ext)).read_bytes()).hexdigest()) for ext in ['pdf','pptx']})
(OUT/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf8')
print(json.dumps(report,indent=2))
