"""Build matching self-study PowerPoint/PDF role guides from reviewed local screens."""
import sys, json, math, re, shutil, zipfile, html
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent / 'training-python'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

OUT = Path('D:/my-business-website/deliverables/Luo-Heritage-Training')
ROOT = Path('D:/LUO LINGUAL SYSTEM/System/VERSION PROJECTS/luolinguaal - Copy')
SHOTS = OUT / 'screenshots'
LOGO = ROOT / 'public/icons/luo-heritage-512.png'
W, H = 960, 540
INK, BROWN, GOLD, CREAM, MUTED, GREEN = '#252723', '#55351e', '#e5ae42', '#f7f3ea', '#62685e', '#245e4d'
for name, file in [('Arial', 'arial.ttf'), ('Arial-Bold', 'arialbd.ttf'), ('Georgia', 'georgia.ttf')]:
    pdfmetrics.registerFont(TTFont(name, 'C:/Windows/Fonts/' + file))

def slide(title, steps, image=None, result=None, tip=None, diagram=None, notes=''):
    return dict(title=title, steps=steps, image=image, result=result, tip=tip, diagram=diagram, notes=notes)

LOGIN = slide('Sign in to your own account', [
    'Open the Luo Heritage address supplied by your organisation. Choose Sign in or Log in.',
    'Enter your email and password, then choose Log in. Check your spelling if it fails.',
    'Community members may also use Google. Staff should use the account supplied by their administrator.'
], '01-sign-in', 'Your account or dashboard opens.', 'Never share your password. Use only an address you trust.')
RESET = slide('Forgot your password?', [
    'On the sign-in screen, choose Forgotten password? Select Reset password.',
    'Enter your account email. Choose Send email, then check your inbox and spam folder.',
    'Open the newest reset link within 30 minutes. A link can be used only once.'
], '03-reset-request', 'An account email arrives if the address is eligible and the email service is available.', 'If no email arrives, contact the person who manages your account. Do not keep requesting links.')
NEW_PASSWORD = slide('Choose a new password', [
    'Open the password reset link from your own inbox.',
    'Enter a new password of at least 12 characters. A short phrase with several words is easier to remember.',
    'Choose Confirm. Return to sign in and use the new password on each device.'
], '04-reset-confirm', 'The password changes and old sign-ins are ended.', 'Do not forward the reset link. If it has expired, request a new one.')
MOBILE = slide('Use the library on a phone', [
    'Open a collection, such as Muma (Dictionary), from the menu.',
    'Use the picture, headphones, video or page icon to choose a type of content. The four-square icon shows all types.',
    'Choose the list or grid button to change the view. Scroll the entries below the controls.'
], '05-mobile-library', 'The collection shows entries matching your chosen filter.', 'An empty result can mean there are no published entries of that type. Try All.')
SOURCE = slide('Check permission before sharing', [
    'Under Source and access, choose Source permission. If unsure, leave Not sure yet / ask a reviewer.',
    'Choose how to credit the source. Use a person’s name only when they agree.',
    'Keep restricted, sacred or research-only material private. Ask a reviewer before changing access.'
], '10-source-permission', 'The draft records who may use the material and how the source should be credited.', 'Permission to collect material is not automatically permission to publish it.')
ATTACH = slide('Add pictures, audio or documents', [
    'Find Media and transcripts. Choose a file under the matching type.',
    'For an oral history, confirm permission before choosing Start recording. Stop when finished.',
    'Check that the correct files are selected. Keep the app open while an upload is in progress.'
], '09-attach-files', 'The selected files stay with the contribution when it is saved successfully.', 'If microphone access is refused, allow it in your device settings or attach an existing recording.')
SAVE = slide('Understand the three save choices', [
    'Save on this device keeps a local copy, including selected files. It does not submit your work.',
    'Save private draft saves an online draft for your account. You need a connection.',
    'Submit for reviewer sends the contribution for checking. Wait for confirmation before leaving.'
], '11-save-submit', 'You know whether the work is on this device, saved online, or sent for review.', 'Clearing browser or app data can remove local drafts. Save again after making changes.')
PRIVACY = slide('Protect your account and other people', [
    'Use your own account. Sign out on a shared computer or phone.',
    'Do not copy private interviews, class results or account lists into public chats or messages.',
    'If a screen shows information you should not see, stop and tell the account administrator.'
], '25-my-account', 'Your work remains linked to the right person and private information stays protected.', 'Send support the screen name and error message, never your password or a reset link.')
HELP = slide('When something does not work', [
    'Check your internet connection. Read the message on screen before trying again.',
    'If your sign-in has expired, sign in again. If a button is missing, ask which role and language you have.',
    'If it still fails, note the time, screen, button and exact message. Send these to your account administrator.'
], '01-sign-in', 'You can explain the problem clearly without exposing private information.', 'Before repeating an upload or submission, check whether the first attempt was already saved.')
FLOW = slide('How a contribution reaches readers', [
    'The uploader or contributor prepares a draft and submits it.',
    'A cultural reviewer checks accuracy. An editor makes requested corrections and returns the item for review.',
    'A publisher checks the final item and permission. Only an approved public release becomes visible to readers.'
], diagram=['Prepare draft', 'Cultural review', 'Publisher check', 'Public release'], result='An item may return for changes. Drafts and private approvals are not public publication.', tip='Different people should perform the independent checks. Do not approve your own work.')
REVIEW = slide('Complete the cultural checks', [
    'Open Cultural Review and choose the item. Read its text, source details and attached files.',
    'Check spelling, translation, grammar, pronunciation, examples and cultural context. Tick only what you have checked.',
    'Add useful notes. If the item is correct, choose Approve & Send to Publisher.'
], '12-cultural-review', 'The item moves to the publisher queue after a successful cultural review.', 'If you cannot check a recording or meaning, ask for help rather than ticking the box.')
EDIT = slide('Correct an item returned by a reviewer', [
    'Open Editor Queue. Find the item and choose Edit details.',
    'Read the requested changes. Correct the text or files, then save your changes.',
    'Return to the queue and choose Send to cultural review. The reviewer must check the corrected version.'
], '14-editor-queue', 'The corrected item returns to cultural review.', 'Saving an edit does not publish it. Keep the source’s meaning and credit.')
PUBLISH = slide('Check carefully before choosing Publish', [
    'Open Publishing Queue and read the culturally reviewed item.',
    'Check the text, attached files, source permission and public access settings. Add a note if needed.',
    'Choose Publish only when the item is ready. Read the confirmation: an internal approval can remain private.'
], '15-publish', 'Eligible public material becomes available to readers; internal material remains protected.', 'A Publish button is not permission to release sacred, restricted or research-only material.')
STUDY = slide('Open a lesson and keep your progress', [
    'Choose Lessons, games & calendar. Open a published lesson.',
    'Read the lesson and answer its questions. Choose Finish & save progress.',
    'Read the result and feedback. Sign in with the same account on another device to keep your learning together.'
], '27-lesson-practice', 'You see a saved result or a message saying the attempt is waiting to synchronize.', 'A score helps you learn; it does not replace a teacher’s review.')
OFFLINE = slide('Prepare to study without internet', [
    'While connected, open a lesson that allows downloads. Choose Download text for offline use.',
    'Keep using the same browser or app profile. Downloaded text can be read without internet for up to seven days.',
    'Reconnect and choose Refresh & sync. Keep the app open until waiting attempts are synchronized.'
], '27-lesson-practice', 'Permitted lesson text is available offline and saved attempts can synchronize later.', 'These downloads contain text, not the audio or video files. Some lessons cannot be downloaded.')
CHAT = slide('Ask the AI chatbot a clear question', [
    'Open Chat or Chatbot. Choose your answer language and read the processing permission notice.',
    'If you agree, allow AI processing. Ask one clear question and wait for the response.',
    'Open any sources provided. Ask a follow-up question if you need a simpler explanation.'
], '24-chat', 'You can read an answer and inspect any linked sources.', 'AI can make mistakes. Do not enter private interviews, passwords or learners’ personal information.')
ACCOUNT = slide('Find your access and saved chats', [
    'Open the account menu and choose My account & saved chats.',
    'Read My access to check your assigned role. Open a chat from Recent or Archived.',
    'For your name, sessions or notifications, open Profile, sessions & notifications from the menu.'
], '25-my-account', 'You can find your own account details and earlier conversations.', 'Choosing an interest such as teacher or researcher does not give staff access.')
ENROLL = slide('Join a learning course', [
    'Open Learning pathways & courses. Choose a published course and read its description.',
    'Open the course and enroll using the button shown. Start with the available lesson.',
    'Use My enrollments to return later. If lessons are locked, complete the earlier lesson and required score first.'
], '28-courses', 'The course appears among your enrollments and shows your lesson progress.', 'If no courses are available, ask your teacher which published course to use.')
CLASS_JOIN = slide('Join a class with your teacher’s code', [
    'Sign in. Open Teaching & my classes and choose My classes.',
    'Ask your teacher for a current code. Enter it under Join a class, then choose Join class.',
    'Open the class and read its assignments. Use your own account for all submissions.'
], '29-classrooms', 'Your class appears in My classes.', 'Class codes expire after seven days and may be replaced earlier. Ask for a new code if needed.')
USERS = slide('Help users without changing the wrong account', [
    'Open Users. Find the person by their own name and email.',
    'Open their details and check their language and current access before changing anything.',
    'Use only the changes allowed by your role. Read every confirmation before suspending or activating accounts.'
], '18-users', 'The intended account has the correct status and access.', 'Do not select every account to solve a problem affecting one person.')
ADMINS = slide('Create a staff account for the right person', [
    'Open Admins, then choose + Add Admin. Confirm the person’s email and name.',
    'Choose only the role and language the person needs. Set a strong temporary password.',
    'Save the account. Share access details privately with that person and ask them to set their own password.'
], '17-add-admin', 'The staff member appears in Administrators with the intended role and language.', 'Only the Master Super Admin can create or remove Super Admin authority.')

decks=[]
def deck(stem,title,subtitle,goal,topics,pages,practice):
    intro=slide('Start here: learn at your own pace', [
        'Read one slide, complete its action, then move to the next slide. Tap means the same as click.',
        'Use your own account. Your organisation supplies the correct website address and any staff access.',
        'Screenshots show the current interface with practice details. Actual names, content and menus may differ.'
    ], diagram=['Read the step', 'Try it yourself', 'Check the result'], result=goal, tip='If your menu looks different, first check My access. Do not ask someone to lend you their account.')
    contents=slide('What this guide covers',topics,diagram=None,result='Use the slide numbers to return to a task. The PDF has page bookmarks; PowerPoint has a contents view.')
    closing=slide('Try it yourself: your readiness check',practice,diagram=['Try', 'Check', 'Ask if unsure'],result='You are ready for routine work when you can complete these steps without help.',tip='Use only approved practice material. Do not publish, delete accounts or change live settings just to practise.')
    decks.append(dict(stem=stem,title=title,subtitle=subtitle,goal=goal,slides=[intro,contents,LOGIN,RESET,NEW_PASSWORD]+pages+[PRIVACY,HELP,closing],sources=[]))

deck('01_Community_Members','Community members','Read, learn, contribute and manage your account',
 'You will be able to find cultural material, ask questions and submit your own contribution.',
 ['Create and recover an account; understand the menu.','Read collections, save items, learn and ask questions.','Prepare a private contribution and send it for review.'],[
 slide('Create a community account',['Choose Create new account from the sign-in screen. Enter your name, email and date of birth.','Choose your language and interests. Read and accept the Terms and Privacy Policy.','Set your password and submit. If approval is required, wait for confirmation before signing in.'],'02-register','Your account is created or clearly marked as awaiting approval.','You can choose Google registration if offered, then complete any requested profile information.'),
 slide('Find your way around',['Open the menu on the left. On a phone, use the menu button first.','Choose a collection such as Muma, Ngero or stories. The smaller label explains its content.','Use the same menu for lessons, courses, saved items, contributions and your account.'],'07-public-menu','You can return to a collection or personal screen without starting again.'),
 MOBILE,
 slide('Read, listen and change the view',['Open an entry or choose the grid button, then select its card.','Use See description or See this in English when offered. Audio and video use their own play controls.','Use All if a filter shows no results. Save item is available in signed-in views that offer it.'],'06-desktop-library','You can read the entry and use the media that has actually been published.','An unavailable English description does not mean the original entry is missing.'),
 slide('Return to saved words and items',['Sign in and choose Save item on an entry or lesson where that button is available.','Open Saved words & items from the menu. Choose Open collection to return to its source.','Choose Remove saved item when you no longer need the bookmark.'],'26-lessons','Your saved list helps you find material again on the same account.','Saving a bookmark does not download a file for offline use.'),
 STUDY,OFFLINE,CHAT,ACCOUNT,
 slide('Start your own contribution',['Choose Contribute & submissions or the Contribute button inside a collection.','Select the collection and choose New private draft. Enter a short title and the requested details.','Use your own material or material you have permission to share. Mark uncertain facts for review.'],'08-contribution-title','A contribution form contains your draft; it is not published.'),ATTACH,SOURCE,SAVE,FLOW,
 slide('Check what happened to your submission',['Open Contribute & submissions, then find My online submissions.','Read the status. If asked to make changes, choose Edit draft or requested revision.','Read the reviewer’s request, make the correction, and submit again when ready.'],'11-save-submit','You can tell the difference between a draft, an item under review and published material.','If an item is locked for review, ask the reviewer to return it rather than creating a duplicate.')
 ],['Find one collection and change its content filter.','Open a lesson or ask a simple AI question, then check any source.','Prepare a private practice draft and identify the option that sends it for review.'])

deck('02_Students','Students and learners','Study lessons, join classes and keep your progress',
 'You will be able to enroll, join a class, study and find your feedback.',
 ['Sign in, recover access and find your learning screens.','Study lessons and courses, including downloaded text.','Join a class, submit work and read teacher feedback.'],[
 MOBILE,ENROLL,STUDY,OFFLINE,CLASS_JOIN,
 slide('Complete a class assignment',['Open your class, then choose Open assignment. Read the instructions and due date.','Choose Study assignment lesson and complete the lesson. Return to the assignment when ready.','Enter the work requested and use the submission button shown. Wait for a saved confirmation.'],'29-classrooms','Your assignment shows your submission rather than an unsaved answer.','If the assignment is closed or the lesson unavailable, contact your teacher; do not submit to a different class.'),
 slide('Read feedback and try again',['Open the same class assignment after your teacher reviews it.','Choose Refresh assignment and feedback. Read the score and the teacher’s comments.','If a new submission is requested and allowed, correct your work and submit it again.'],'27-lesson-practice','You know what to improve and which result belongs to you.','Your quiz score and your teacher’s assignment grade can differ.'),
 CHAT,ACCOUNT,
 slide('Keep your own learning record',['Use the same account on your phone and computer.','Before changing devices, reconnect and use Refresh & sync for waiting attempts.','Do not leave a class just to change devices: leaving can delete your work in that class.'],'25-my-account','Your work stays linked to the intended account.','Tell your teacher if you accidentally created two accounts.')
 ],['Enroll in the course your teacher names and open its first available lesson.','Join your class with a current code and find the assignment instructions.','Save one practice result and show where you would read teacher feedback.'])

deck('03_Teachers_and_Lecturers','Teachers and lecturers','Guide classes, set assignments and review learning',
 'You will be able to run a class and help students find their work and feedback.',
 ['Get teaching access and create or open a class.','Invite learners, set assignments and review submissions.','Keep class results private and close a class responsibly.'],[
 slide('Check that you have teaching access',['Sign in with your own account. Open Teaching & my classes.','Look for Create class or an existing class where your role is teacher.','If these are missing, ask the language administrator to assign teaching access. Selecting teacher as an interest does not grant it.'],'29-classrooms','You know whether you can teach a class or are viewing as a learner.'),
 slide('Create a teaching class',['Open My classes and choose Create class.','Enter the class title and requested details. Select a published course if the class needs one.','Choose Save teaching class. Keep the new class join code private until you invite learners.'],'29-classrooms','The class opens and you can see its invitation controls.','Teaching a class does not automatically let you publish cultural content or courses.'),
 slide('Invite your learners safely',['Open the class and find Class invitations. Share its current join code with the intended learners.','Ask learners to sign in and use Join a class. Check the member list afterwards.','If a code is shared too widely, choose Replace class join code or Disable class join code.'],'29-classrooms','The correct learners are listed in the class.','Codes expire after seven days. A replacement invalidates the previous code.'),
 slide('Set a useful assignment',['Open the class and choose Create assignment.','Choose an available published lesson. Write a clear title, instructions and due date.','Choose Save class assignment. Open it again to check exactly what learners will see.'],'29-classrooms','The assignment is listed for members of this class.','Use a lesson in the correct language. If a lesson is withdrawn, choose an approved replacement.'),
 slide('Review a learner’s submission',['Open the class assignment and choose the learner’s submission.','Read their work, enter the score and write feedback that explains one clear next step.','Choose Save teacher grade and feedback. Ask for another submission when correction is needed.'],'29-classrooms','The learner can refresh the assignment to see the saved grade and feedback.','Check the learner’s name before saving. Do not discuss individual results in public chats.'),
 slide('See the class gradebook',['Open the class and choose View gradebook.','Read the assignment columns and learner rows. A teacher’s grade takes priority over a synchronized lesson score.','Use Export gradebook CSV only if you are allowed to keep a spreadsheet copy. Store it privately.'],'29-classrooms','You can identify submitted, ungraded and completed work.','A downloaded gradebook contains personal information. Do not share it as a public attachment.'),
 slide('Manage changes without losing work',['Use Edit assignment before closing submissions. Read any warning when changing instructions.','Choose Close submissions when no more work should be accepted.','Remove a class member only after checking the consequences: that action can remove their class submissions.'],'29-classrooms','Class membership and assignment availability match your intention.'),
 slide('Share teaching responsibility',['Open the class and find Assign a teacher by existing account email.','Check the colleague’s exact account email and grant teaching access only if you have permission.','Confirm their role in the member list. Keep at least one teacher in each active class.'],'29-classrooms','The intended colleague has teaching access to this class.','A class teacher is different from a system administrator.'),
 ENROLL,STUDY,OFFLINE,
 slide('Finish a term or close a class',['Save any permitted reports before closing the class.','Choose Archive class only when teaching, invitations and submissions should end. Read the confirmation.','Explain to learners that archived classes remain readable to existing members, while new activity is closed.'],'29-classrooms','The class is archived rather than accidentally deleted.', 'Do not leave or remove the last teacher from an active class.'),
 CHAT
 ],['Open a practice class and explain how a student joins it.','Identify Create assignment, the feedback area and View gradebook.','Explain the difference between closing submissions, archiving a class and removing a learner.'])

deck('04_Uploaders','Uploaders','Prepare accurate drafts and send them for review',
 'You will be able to add the right material, record permission and follow its review status.',
 ['Confirm your assigned language and find Upload Content.','Prepare text, attachments, source details and private drafts.','Submit once and follow any requested corrections.'],[
 slide('Open your uploading workspace',['Sign in using the account supplied to you. Open the admin dashboard.','Choose Upload Content. Work only in the language assigned to you.','If no collection can be created, check your assigned language with the administrator.'],'31-language-admin-menu','Your content list opens for the languages you are allowed to upload.','This screenshot shows a language administrator’s wider menu. Uploaders normally see Dashboard and Upload Content.'),
 FLOW,
 slide('Prepare before you upload',['Check the collection, language, title and file names before starting.','Confirm who supplied the material and what permission was given. Keep evidence through the organisation’s approved process.','Use a draft when a meaning, translation or permission is uncertain. Do not guess.'],'08-contribution-title','You have the information needed for a useful draft.'),
 slide('Enter the contribution details',['Open the appropriate new-content form from Upload Content.','Enter the title and all required collection fields. Read the text again for missing words.','Add a clear description and choose Save draft. An admin upload is a draft until submitted for review.'],'08-contribution-title','The online draft appears in your content list.','Admin forms use Save draft or Save changes. Community forms use Save private draft and Submit for reviewer.'),
 ATTACH,SOURCE,
 slide('Add collection details when required',['Open the extra source or collection details section.','Record the known collection date, source reference, location and collector. Leave unknown details for review.','Do not put personal phone numbers, sacred site coordinates or private participant details in a public description.'],'10-source-permission','Reviewers can understand where the material came from without exposing private details.'),
 SAVE,
 slide('Send the online draft for cultural review',['Return to Upload Content and find the saved draft.','Use the available submit action. Confirm that you are submitting the intended item once.','Refresh the list and check its new status. Wait for review; you cannot publish it as an uploader.'],'11-save-submit','The draft is marked as submitted or under review.','The screenshot shows community save choices. In the admin list, use the submit action for the saved draft.'),
 slide('Respond to a correction request',['Find the returned item and read the reviewer’s notes.','Open the editable draft, make only the necessary corrections and save.','Submit the corrected item again or follow the editor’s handover. Tell the reviewer what changed.'],'14-editor-queue','The corrected version is ready for another independent check.','Editors use Editor Queue. Uploaders use the editable items available in Upload Content.'),
 slide('Check files before you finish',['Reopen the saved item and inspect every attached file.','If an upload failed, reconnect and check whether the text or file was already saved before retrying.','Keep your original source file until the organisation confirms it is safely stored.'],'09-attach-files','The item has the intended files and no accidental duplicate uploads.')
 ],['Locate your assigned collection and create a private practice draft.','Explain the permission setting you would choose when permission is uncertain.','Show where you would submit the draft, and explain why you cannot publish it yourself.'])

deck('05_Cultural_Reviewers','Cultural reviewers','Check accuracy, meaning and respectful use',
 'You will be able to make a clear, independent cultural review decision.',
 ['Open Cultural Review and examine the item and its source.','Approve accurate material or request clear changes.','Use independent voting for a proposed cultural rejection.'],[
 FLOW,
 slide('Open the right review item',['Choose Cultural Review in your staff menu.','Open a submitted item in your assigned language. Check that you did not prepare the material yourself.','Read its source, access settings, full text and attached media before deciding.'],'12-cultural-review','You are reviewing the current version of the intended item.','An item can change while another reviewer works on it. Refresh if the system asks you to.'),
 REVIEW,SOURCE,
 slide('Request an editor’s correction',['Describe the problem in the notes: say where it is, why it matters and what needs checking.','Choose Request Editor Changes when a correction is needed.','The editor fixes the draft and sends it back to cultural review. Check the corrected version again.'],'12-cultural-review','The editor receives a clear change request; the item is not published.','Write useful comments such as “Check the meaning in paragraph two,” rather than “Wrong.”'),
 slide('Raise a serious cultural concern',['If material may be culturally false or harmful, explain your concern in the notes.','Choose Propose rejection for reviewer vote. This asks other assigned reviewers to check it.','Do not treat your proposal as a completed rejection. Wait for the required independent votes.'],'12-cultural-review','The concern goes to an independent review vote.','For a serious privacy problem, also contact the responsible administrator promptly.'),
 slide('Vote on another reviewer’s proposal',['Open the proposed cultural rejection. Read the reason and the original source.','If you are eligible, choose whether the item is culturally valid or invalid. Give a clear reason for an invalid vote.','Check the confirmation and vote count. You cannot vote on your own proposal or vote twice.'],'13-rejection-vote','Your independent vote is recorded. The system applies the required threshold.'),
 slide('Keep a useful review history',['Open Review History to find your recorded actions.','Use the item’s title and review notes when explaining a decision to an editor or coordinator.','If an error is found later, report it with the item reference rather than editing history or concealing the earlier decision.'],'12-cultural-review','Other staff can understand how the decision was reached.'),
 slide('Review learning material with the same care',['Open Lessons, games & calendar if your role offers the content workspace.','Read a submitted lesson, its questions, answers, source and audience. Use the review action only if independently assigned.','A course review is separate from reviewing its individual lessons. Ask the coordinator if the action is not available.'],'26-lessons','Learning material receives the appropriate independent checks.','A role that can review material does not automatically let its author approve their own work.')
 ],['Explain one checklist item using an approved practice contribution.','Describe how to request an editor change and how the item returns to you.','Explain why a proposed rejection needs independent votes.'])

deck('06_Content_Editors','Content editors','Make clear corrections and return work for review',
 'You will be able to correct content without bypassing the review process.',
 ['Find returned items in Editor Queue.','Correct text, translations and files while protecting the source.','Save the changed version and return it for cultural review.'],[
 FLOW,EDIT,
 slide('Make a careful edit',['Read the reviewer’s request before changing the draft. Keep the original meaning.','Correct spelling, formatting, descriptions and the requested collection details. Do not invent missing translations.','Save changes, reopen the item and check that the corrected version is shown.'],'08-contribution-title','The saved draft reflects the requested correction.','For a disputed meaning, ask a cultural reviewer rather than silently replacing it.'),
 ATTACH,SOURCE,
 slide('Return the item to the reviewer',['Go back to Editor Queue after saving your changes.','Choose Send to cultural review beside the intended item.','Read the success message. The item should leave the editable queue when the handover succeeds.'],'14-editor-queue','The current corrected version is ready for cultural review.','If the item has already moved, refresh instead of sending a stale copy.'),
 slide('Prepare a lesson or game draft',['Open Lessons, games & calendar and use Content workspace when offered.','Choose Create content. Write the title, content, level and audience; add checked questions and answers.','Save the private draft and submit for independent review. Do not approve your own resource.'],'26-lessons','The resource is a draft or submitted item, not automatically public.','A written question game is different from a live multiplayer game.'),
 slide('Build a course from approved lessons',['Open Learning pathways & courses and enable Course staff workspace.','Choose Create course. Use published lessons from the same language, place them in order and set required scores.','Save course draft, open it and submit it for independent course review.'],'28-courses','The course draft links the intended approved lessons.','Changing a published course requires another review before publication.'),
 slide('Avoid duplicate or unfinished corrections',['If saving fails, keep the draft open and check the connection.','Reopen or refresh only after preserving unsaved text in a trusted place. Check the current item version.','If another person has already changed the item, read their changes before trying again.'],'14-editor-queue','You update the intended current version and avoid overwriting another person’s work.')
 ],['Find a practice correction and identify the requested change.','Save a corrected draft and locate Send to cultural review.','Explain why editing, cultural review and publication are separate actions.'])

deck('07_Publishers','Publishers','Make the final release decision responsibly',
 'You will be able to check a reviewed item and distinguish public release from internal approval.',
 ['Find culturally reviewed items in Publishing Queue.','Check source permission, access and media before release.','Publish eligible material and verify what readers can see.'],[
 FLOW,PUBLISH,SOURCE,
 slide('Use a final pre-publication check',['Confirm the title, language and current version are correct and the cultural review is complete.','Open the files and read descriptions, translations and source credit. Check that public sharing is actually permitted.','If anything is unclear, stop and contact the reviewer or language administrator before publishing.'],'15-publish','The release decision is based on the complete item and its permission evidence.'),
 slide('Understand the confirmation message',['After choosing Publish, read the full success message.','Public release means eligible readers can see the item. Internal curation means it remains private.','Do not change consent or access merely to make an internally approved item appear in the public library.'],'15-publish','You know whether the item was publicly released or privately curated.'),
 slide('Check the reader’s view',['Open the public library, preferably in a separate signed-out browser window.','Choose the correct collection and content type. Find the released item and try its files.','If it is missing, confirm the release message, language, permission and selected filter before escalating.'],'06-desktop-library','The intended public item is readable and its files are usable.','Do not test private content by sharing a staff-only link with a member of the public.'),
 slide('Keep the publication record',['Open Published History to find earlier releases.','Record any concern with the item’s title and reference, and notify the language administrator.','If an urgent rights or privacy issue is found, request the approved withdrawal process. Do not assume a saved edit removes a public item.'],'15-publish','Publication decisions and later concerns can be traced.'),
 slide('Publish lessons and courses separately',['In the learning content workspace, inspect validated resources and use publish only when permitted.','For a course, open its staff view and choose Publish course after independent review.','Confirm that both the course and its linked lessons are available to the intended learners.'],'28-courses','Learners see a usable course with available lessons.','You cannot independently approve a resource you authored yourself.')
 ],['Explain the checks you make before Publish.','Distinguish a public release from an internal approval.','Find the public collection and explain how you would report an incorrect release.'])

deck('08_Language_Administrators','Language administrators','Coordinate people and content for your language',
 'You will be able to manage your language’s work while keeping staff duties clear.',
 ['Understand your language’s menu, people and content.','Coordinate uploads, cultural review, editing and publication.','Check teaching access, translations, activity and problems.'],[
 slide('Read your language workspace',['Open the admin dashboard and confirm the selected language.','Use Content, Cultural Review, Editor Queue and Publishing Queue to find work at each stage.','Use Media, Users and Transcripts only for the material and people in your permitted scope.'],'31-language-admin-menu','You can direct a question to the right screen and staff member.'),FLOW,USERS,
 slide('Give a person only the access they need',['Open the person’s account details and verify their email.','Use the role controls available to you for your language. If staff creation is not offered, ask a Super Admin.','Confirm the assignment in the person’s access list and ask them to sign in again if the menu has not refreshed.'],'18-users','The right person can do their assigned task in the intended language.','Language access does not allow you to grant Super Admin ownership.'),
 slide('Coordinate a review without skipping stages',['Check the oldest waiting items and who is assigned to each stage.','Ask the uploader to complete missing source details, the reviewer to check accuracy, or the editor to make requested changes.','Keep independent review: broad access does not make it appropriate to approve your own submission.'],'12-cultural-review','Work moves through the intended people and checks.'),REVIEW,EDIT,PUBLISH,
 slide('Check media and translation quality',['Open Media and the relevant content item. Check files, labels and access settings.','Use Translations to review the language labels and available translations. Missing text needs a checked translation.','Ask the responsible content editor or reviewer to correct inaccurate material.'],'09-attach-files','The collection is usable in the intended language without misleading labels.'),
 slide('Handle teaching access and learning content',['Confirm that a teacher has been given actual teaching access; an interest on their profile is not enough.','Check that lessons and courses are published only after independent review.','Ask teachers to handle class membership, assignments and gradebooks within their classes.'],'29-classrooms','Teachers can use the right classes and learners can open approved lessons.'),
 slide('Read activity and AI reports carefully',['Open Analytics for the activity recorded by the system.','Use AI Monitoring to review recorded questions, feedback and concerns within your allowed access.','Treat low counts as a reason to investigate usage, not as proof of a technical fault. Escalate private-data or accuracy concerns.'],'31-language-admin-menu','You can explain recorded activity without inventing missing results.'),
 slide('Handle a permission or privacy complaint',['Note the affected item and what the person says is wrong. Avoid copying the private content into your report.','Notify the responsible institutional owner and arrange the permitted restriction or withdrawal action.','Check that the final public view matches the decision and retain the review record.'],'10-source-permission','A concern has an owner and a documented response.','Do not resolve a complaint by silently removing source evidence or review history.')
 ],['Name the person responsible for each stage of a contribution.','Check a practice account’s language and duties without granting extra power.','Explain how you would handle a missing teacher button or a privacy complaint.'])

deck('09_Super_Admins','Super Admins','Oversee assigned languages and support staff',
 'You will be able to coordinate access and monitor work without confusing oversight with editorial permission.',
 ['Understand what your wider menu allows and what remains restricted.','Manage staff for your assigned languages and help users.','Monitor content, activity and system settings responsibly.'],[
 slide('Understand your wider menu',['Use Workflow and Library to inspect content, people and collections.','Use Insights to see Analytics, AI Monitoring and Audit Logs.','Administration contains Languages, Admins, Backup and System. Visibility does not mean every action is allowed.'],'32-super-admin-menu','You can find the right area without assuming that every visible button is yours to use.'),
 slide('Know the limits of Super Admin access',['You oversee the languages assigned to you and may create permitted staff within that scope.','Editing, cultural review, uploading and publishing still require their assigned language duties.','Only the Master Super Admin can create or remove Super Admin authority. Database backups are also reserved for that owner.'],'16-administrators','You know when to act and when to refer the request to the Master Super Admin.'),
 ADMINS,USERS,FLOW,
 slide('Check staff assignments regularly',['Open Admins and verify each person’s email, status, language and duty.','If a person changes duties, update access through the permitted controls; do not share another account.','Read every revoke or deactivate confirmation. Revoking staff duties is different from deleting the person’s account.'],'16-administrators','Staff have the access needed for current work and no unnecessary authority.'),
 slide('Help a user regain access',['Confirm which email and language they should be using. Do not ask for their password.','Guide them to Forgotten password? and the newest email link. Check whether the account is active.','If password emails are not being delivered, ask the authorised system owner to check Email settings and the test result.'],'03-reset-request','The user receives the right recovery instructions without exposing their credentials.'),
 slide('Review activity and recorded decisions',['Use Analytics and AI Monitoring to inspect recorded activity and feedback.','Use Audit Logs to see who changed a record or setting and when.','Investigate unexpected changes with the responsible staff member. Do not alter evidence to hide a mistake.'],'32-super-admin-menu','You can follow a concern to the relevant person and recorded action.'),
 slide('Use System settings carefully',['Open System and read the selected tab before changing anything.','Change only a setting you understand and are authorised to manage. Save the intended group of settings.','Check the result afterwards. For service keys, backup, legal ownership or an unavailable control, contact the Master Super Admin.'],'19-system','A deliberate saved change has the intended effect.','Do not include secret service keys in screenshots, slides, email or support messages.'),
 slide('Coordinate content and learning concerns',['Route incomplete uploads to uploaders, corrections to editors and cultural concerns to reviewers.','Ask the assigned publisher to check release readiness. Do not bypass the independent checks.','For class work, involve the teacher. Keep learner results and research material private.'],'12-cultural-review','The right person handles the issue without excessive access.')
 ],['Explain the difference between a Super Admin and the Master Super Admin.','Create no live accounts during practice; instead identify the correct role and language in the form.','Find the recovery, monitoring and escalation paths for a reported problem.'])

deck('10_Master_Super_Admin','Master Super Admin','Own platform access and protect the institutional system',
 'You will be able to oversee high-impact actions, delegate responsibly and keep the platform’s records safe.',
 ['Manage Super Admin authority, languages and institutional settings.','Check account recovery, backups, public terms and service health.','Use documented checks before high-impact changes and handover.'],[
 slide('Your role carries overall responsibility',['Use your own protected account. Confirm My access shows Master Super Admin.','Keep the organisation’s ownership, support arrangements and recovery contact current.','Delegate ordinary work to named staff; broad access is not a reason to bypass independent review.'],'25-my-account','You know which responsibilities must remain with the institutional owner.','The screenshot uses a community practice account. Your own My access must show Master Super Admin.'),
 slide('Create or change Super Admin authority',['Open Admins, then + Add Admin to prepare a new staff account.','Select Super Admin only for a trusted person with an agreed responsibility. Assign the languages they oversee.','Check the saved account carefully. Read the consequences before revoking or deactivating a Super Admin.'],'17-add-admin','Only the intended person receives the intended authority.','Never create a shared administrator account for a team.'),
 ADMINS,
 slide('Maintain your staff list',['Open Administrators and review each name, email, language and role.','Remove unneeded authority through the permitted controls after confirming ownership of ongoing work.','Keep cultural reviewers, editors and publishers distinct wherever possible. Record the reason for important access changes.'],'16-administrators','Current staff responsibilities are clear and former staff cannot retain unwanted access.'),
 slide('Add a language deliberately',['Open Languages and choose + Add Language.','Follow the numbered language setup: enter the code, name and local name, then check the module and administrator choices.','Review the summary before saving. Creating a language does not create reviewed translations or public cultural content.'],'33-new-language','The intended language exists with its configuration ready for checked content.'),
 slide('Review a language before activation',['Open Languages and select the language.','Check its labels, assigned staff and collections. Confirm that any public content is reviewed and permitted.','Use the available activation controls only when the organisation agrees it is ready.'],'21-languages','The public language selection reflects your deliberate decision.'),
 FLOW,
 slide('Check account email settings',['Open System, then Email. Confirm the approved provider, verified sender address and private service key.','Save the email settings. Choose Send test email to send to your own signed-in account.','Check both the screen result and your inbox. A provider accepting a message is not proof it reached the inbox.'],'20-email-settings','You have evidence that the account email path is configured and can be tested.','Keep keys secret. Password reset cannot send email while provider credentials or sender approval are missing.'),
 slide('Help with a failed password reset',['Confirm Email settings and the test result before asking the user to retry.','Ask the user to request a fresh link, check spam and use the newest link within 30 minutes.','Keep the user’s reset link and new password private. Record the problem without including either.'],'03-reset-request','The user can follow a safe recovery path once email delivery is available.'),
 slide('Make and protect a database backup',['Open Backup and choose Backup Now when a backup is authorised.','Wait for confirmation, then check the new file in the list. Store any downloaded copy in the organisation’s protected location.','Ask the technical maintainer to verify recovery. A database backup does not contain the actual picture, audio or video files.'],'22-backups','A protected database snapshot is recorded; separate media backup and recovery checks are arranged.','Backup files contain sensitive account data. Never attach them to a public message or training exercise.'),
 slide('Check public terms and support information',['Open System and review the General settings and Terms & Privacy tab.','Replace draft organisation details with the verified operator and support information before publishing those texts.','Have the responsible institution approve the wording. Check the public pages after saving.'],'19-system','Readers can find the correct organisation and support information.','Do not publish a legal draft containing bracketed placeholders or unapproved promises.'),
 slide('Treat AI settings and feedback responsibly',['Use API Keys and AI Monitoring only for the services and records you are permitted to manage.','Keep service credentials private. Test ordinary questions without submitting restricted interviews or personal information.','Investigate wrong answers using their sources and feedback. Do not describe untested language or image models as reliable.'],'24-chat','AI support is used with clear source and privacy checks.'),
 slide('Use maintenance and ownership controls carefully',['Before enabling maintenance, tell the affected staff and agree when access will return.','Read each control and confirmation. Confirm the public view afterwards.','Do not delete the Master Super Admin account to transfer responsibility. Arrange the approved institutional ownership process first.'],'19-system','High-impact changes have an owner, a reason and a checked result.'),
 slide('Prepare an orderly handover',['Give each person their own account, role guide and approved website address.','Keep private service details, backup locations and recovery contacts in a restricted handover record.','Record which service and device checks have actually passed. Treat physical-server migration and store release as separate planned work.'],'32-super-admin-menu','The next responsible person can operate the system without sharing passwords or guessing what was tested.')
 ],['Identify which duties you may delegate and which require Master Super Admin authority.','Explain how to check account email delivery without sharing a key or reset link.','Explain what the database backup includes, what it excludes, and how you would request a recovery check.'])

deck('11_Researchers','Researchers and institutional readers','Find usable sources and respect collection permissions',
 'You will be able to distinguish public reading from authorised research access and request help safely.',
 ['Use public collections and sources responsibly.','Understand what the researcher interest does and does not allow.','Read internal transcripts only when separately authorised.'],[
 slide('Understand your access before you begin',['A researcher interest describes your purpose. It does not unlock private interviews or staff screens.','Start with the public library and approved sources available to ordinary accounts.','For internal research material, contact the institutional custodian. Use only the access they formally assign.'],'25-my-account','You know which material you may use without borrowing another person’s account.','This guide does not promise a self-service private-dataset download portal.'),
 MOBILE,
 slide('Use a source, not only an AI answer',['Open the approved entry and read its description, language and source credit.','For an AI answer, open the cited platform entries and compare the answer with them.','Record the source title and permitted reference details for your research notes. Follow the institution’s citation and reuse rules.'],'06-desktop-library','Your notes refer to actual approved material rather than an uncited generated answer.'),CHAT,
 slide('Request private research access properly',['Describe your research purpose, the collection needed and who will use it.','Ask the institutional custodian what permission, consent and access conditions apply.','Wait for approval. Do not treat a visible title, an imported file or a colleague’s copy as permission to reuse it.'],'23-transcripts','The access request can be considered by the responsible person.','Research-only, teaching and public sharing are different permissions.'),
 slide('For authorised reviewers: open transcripts',['This step applies only if you have separately been assigned a role that can read the internal transcript library.','Open Transcripts from the staff menu. Choose the permitted session in your assigned language.','Check the session reference and read only the material allowed by your approval.'],'23-transcripts','An authorised reader reaches the correct internal session.','If Transcripts is missing or access is refused, ask the custodian. Do not try another user’s login.'),
 slide('Read transcript sections carefully',['Open the permitted session and choose a segment, or Expand all segments.','Keep the source reference with your notes. An AI summary is not a replacement for the original text.','Do not translate, publish, export or submit internal text to an external AI service without the required permission.'],'34-transcript-segments','Your work retains its connection to the permitted original source.'),
 SOURCE,
 slide('Request an export rather than assuming permission',['Tell the custodian exactly which records and fields are needed and why.','Ask for the approved export and any citation, storage, sharing or deletion conditions.','Check the files and permission record together. Keep participant identities and sensitive locations protected.'],'23-transcripts','Any research copy has a documented purpose and agreed conditions.','An administrator’s technical ability to export is not the same as permission for your study.'),
 slide('Report a source or consent concern',['Note the item title and reference, and describe the concern without copying restricted text.','Contact the institutional custodian or the administrator who granted your access.','Pause any disputed reuse until the responsible person confirms what is permitted.'],'10-source-permission','The rights concern is handled without spreading the disputed material.')
 ],['Find a public source and explain how you would cite it.','Explain why the researcher interest does not grant internal access.','Describe how you would request a transcript or dataset export and check its reuse conditions.'])

START = dict(stem='00_START_HERE',title='Choose your training guide',subtitle='Self-study slide library • eleven roles',goal='Open the guide that matches the work you have actually been assigned.',slides=[
 slide('Welcome to your self-study slide library',['Each role has a PowerPoint slide deck and a matching PDF. Open the PDF for easy reading; use PowerPoint for presenting or editing.','Read one task, try it on your own account, then check the result described on the slide.','Use the guide for your assigned duty. If you have several duties, read each relevant guide.'],diagram=['Choose a role','Read the steps','Try and check'],result='No classroom trainer is needed to follow the guides.'),
 slide('Start with your everyday work',['Community members: browsing, accounts, contributions and AI chat.','Students: courses, classes, assignments and progress.','Teachers and lecturers: class setup, assignments, feedback and gradebooks.'],result='Open guide 01, 02 or 03.'),
 slide('If you prepare or review content',['Uploaders: draft, attach files, record permission and submit.','Cultural reviewers: check accuracy, request corrections and vote independently.','Content editors: correct returned work. Publishers: check the final release.'],diagram=['04 Upload','05 Review','06 Edit','07 Publish'],result='Open the guide matching your assigned content duty.'),
 slide('If you manage the platform',['Language administrators: coordinate people and content in an assigned language.','Super Admins: oversee assigned languages and staff, within their authority.','Master Super Admin: institutional ownership, Super Admin authority and protected system operations.'],result='Open guide 08, 09 or 10. A wider menu does not grant every action.'),
 slide('If you use the system for research',['Open guide 11 for public sources, internal-access requests and transcript guidance.','A researcher interest on your profile is not permission to read private interviews.','Ask the institutional custodian before exporting or reusing restricted material.'],'23-transcripts','Research follows the source’s access and reuse conditions.'),
 slide('About the screenshots',['Screens are captured from the current application components using sample names and practice records.','They are illustrations of the interface, not evidence that an email was delivered, a record was published or a live account was created.','Your menu and records depend on your assigned access. Phone and computer layouts may place controls differently.'],'05-mobile-library','You can recognise controls without mistaking practice details for real cultural evidence.'),
 slide('Share the right files',['Send a person their own role’s PDF or PowerPoint. PDF copies keep the layout consistent.','The role decks contain their own account and recovery instructions, so they can be read independently.','Keep the editable sources for future updates. Review the screenshots again when the application changes.'],result='Distribute the guides as slides, not as a long technical manual.'),
 slide('Need help while studying?',['For a missing button, check your assigned role and language with your account administrator.','For a failed action, note the screen, button, time and exact message. Never send your password.','Practise high-impact tasks only in an approved practice environment. Do not publish or delete real data to test a guide.'],'01-sign-in','Your administrator can help with a clear description of the problem.')
])

# Use the exact role and task views rather than a broader menu as an example.
SCREEN_UPDATES = {
    'Return to saved words and items': '37-saved-items',
    'Complete a class assignment': '49-student-assignment',
    'Read feedback and try again': '52-student-feedback',
    'Create a teaching class': '38-create-class',
    'Invite your learners safely': '39-class-invitations',
    'Set a useful assignment': '42-assignment-editor',
    'Review a learner’s submission': '50-teacher-grading',
    'See the class gradebook': '41-gradebook',
    'Manage changes without losing work': '40-class-assignments',
    'Share teaching responsibility': '39-class-invitations',
    'Finish a term or close a class': '40-class-assignments',
    'Open your uploading workspace': '35-uploader-menu',
    'Enter the contribution details': '45-admin-draft',
    'Send the online draft for cultural review': '54-uploader-draft-list',
    'Respond to a correction request': '08-contribution-title',
    'Prepare a lesson or game draft': '44-lesson-editor',
    'Your role carries overall responsibility': '36-master-access',
    'Know the limits of Super Admin access': '51-super-admin-staff',
    'Check staff assignments regularly': '51-super-admin-staff',
}
PERMISSION_LABELS = slide('What the permission labels mean',[
    'Source permission: has the source agreed to the proposed use? Choose Not sure yet when you do not know.',
    'Consent scope: what use was agreed? Public excerpt means an approved public part; research only and teaching have different limits.',
    'Restriction level: who may see it? Internal and restricted are not public. Sacred material needs the custodian’s special care.'
], '10-source-permission', 'You understand the labels before you change them.', 'If the permission record is unclear, keep the item private and ask the responsible reviewer.')
SEARCH = slide('Find a word in the public dictionary',[
    'On the public website, open the dictionary page and find its search box.',
    'Enter a Dholuo word or an English or Kiswahili meaning. Wait for the results, then read the relevant entry.',
    'If there is no result, check spelling or try another word. A missing published entry is not proof that the word does not exist.'
], '53-dictionary-search', 'You can use the dictionary search where it is offered.', 'The shared app’s Muma screen uses collection filters; it may not show this website search box.')
for guide in decks:
    updated=[]
    for s in guide['slides']:
        if s['title'] in SCREEN_UPDATES:
            s = dict(s, image=SCREEN_UPDATES[s['title']])
        if s['title']=='Open your uploading workspace': s=dict(s,tip='Uploaders normally see Dashboard and Upload Content. Other duties require a separate assignment.')
        if s['title']=='Send the online draft for cultural review':
            s=dict(s,steps=['Return to Upload Content and find the saved draft.','Choose Send to cultural review beside the intended item. Send it only once.','Refresh the list and check its status. You cannot publish an item using only the uploader role.'],tip='The English version link opens a separate translation. Complete it where required before review.')
        if s['title']=='Your role carries overall responsibility': s=dict(s,tip='Check My access on your own account; the screenshot uses a labelled practice account.')
        updated.append(s)
        if s['title']==SOURCE['title']: updated.append(PERMISSION_LABELS)
        if s['title']==MOBILE['title'] and guide['stem'] in ['01_Community_Members','11_Researchers']:updated.append(SEARCH)
    guide['slides']=updated
    if guide['stem']=='04_Uploaders':
        guide['slides']=[slide('Know where your draft is saved',[
            'In Upload Content, Save draft or Save changes saves the item online. Wait for confirmation and check the list.',
            'The community Contribute & submissions form also offers device copies. That is a separate route from the admin upload form.',
            'A device copy is never submitted automatically. To send an admin draft, use Send to cultural review in the online list.'
        ],'45-admin-draft','You know which copy is saved and which action sends it for review.','Keep the original file until the online item and its attachments have been checked.') if s['title']==SAVE['title'] else s for s in guide['slides']]

PROFILE=slide('Update your profile and read notices',[
    'Open Profile, sessions & notifications from the menu.',
    'Change your name or interests and choose Save profile. Interests describe you; they do not grant staff duties.',
    'Read your notifications and use Mark read when offered. Check Device sessions if you need to end an old app sign-in.'
], '30-account-settings', 'Your profile changes are saved and notices are easy to find.', 'Sign out all devices ends your current sign-in too. Sign in again afterwards if you use it.')
VERIFY=slide('Verify your email address',[
    'Open account verification and recovery from your account settings.',
    'Choose Verify email under Action, enter your own account email and choose Send email.',
    'Open the newest email link within 30 minutes and choose Confirm. Use a new link if the earlier one has expired.'
], '03-reset-request', 'The confirmation verifies the address for the intended account.', 'The screenshot shows the same form set to Reset password. Change Action to Verify email for this task.')
DELETE=slide('If you decide to delete your account',[
    'Read the warning first. Deletion removes your sign-in, profile, saved chats and personal progress.',
    'In account recovery, choose Delete account. Request the email link only if you intend to continue.',
    'Open the link and read the confirmation before choosing Confirm. Required cultural-source and review records can be retained.'
], '03-reset-request', 'You understand the consequences before requesting deletion.', 'This is optional, not a practice task. Active class teachers and the Master Super Admin may need to resolve their responsibilities first.')
for guide in decks:
    if guide['stem']=='01_Community_Members':
        position=next(i for i,s in enumerate(guide['slides']) if s['title']==ACCOUNT['title'])+1
        guide['slides'][position:position]=[PROFILE,VERIFY,DELETE]
    for s in guide['slides']:
        # Define the uncommon role word in everyday language instead of
        # expecting independent readers to know repository terminology.
        s['steps']=[t.replace('institutional custodian','person responsible for the collection').replace('institution’s citation','organisation’s citation').replace('custodian’s special care','collection owner’s special care').replace('custodian','collection owner') for t in s['steps']]
        if s['title']=='Read feedback and try again':s['tip']='Sending revised written work clears its previous teacher grade so the teacher can review it again.'
        if s['title']==OFFLINE['title']:
            s['steps'][2]='Reconnect and choose Refresh & sync. Keep the app open until waiting attempts have been saved online.'
        if s['title']==STUDY['title']:s['result']='You see a saved result or a message saying the attempt is waiting to go online.'
        if s['title']=='Start here: learn at your own pace':s['steps'][0]='Read a slide, try its action, then check the result. Tap means click. Gold circles mark important controls in the screenshots.'

def rgb(value): return RGBColor.from_string(value.lstrip('#'))
def wrap(text, font, size, width):
    lines=[]
    for para in text.split('\n'):
        line=''
        for word in para.split():
            proposal=(line+' '+word).strip()
            if pdfmetrics.stringWidth(proposal,font,size)>width and line:
                lines.append(line);line=word
            else: line=proposal
        lines.append(line)
    return lines

class SlideDraw:
    def __init__(self, pres, pdf):
        self.p=pres.slides.add_slide(pres.slide_layouts[6]);self.c=pdf
        self.p.background.fill.solid();self.p.background.fill.fore_color.rgb=rgb(CREAM)
        self.c.setFillColor(CREAM);self.c.rect(0,0,W,H,fill=1,stroke=0)
    def rect(self,x,y,w,h,color,radius=0,line=None):
        shape=self.p.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));shape.fill.solid();shape.fill.fore_color.rgb=rgb(color)
        if radius:
            try: shape.adjustments[0]=.08
            except Exception: pass
        if line:shape.line.color.rgb=rgb(line)
        else:shape.line.fill.background()
        self.c.setFillColor(color);self.c.setStrokeColor(line or color)
        if radius:self.c.roundRect(x,H-y-h,w,h,radius,fill=1,stroke=int(bool(line)))
        else:self.c.rect(x,H-y-h,w,h,fill=1,stroke=int(bool(line)))
    def text(self,text,x,y,w,size=18,color=INK,font='Arial',bold=False,leading=None,link=None):
        font='Arial-Bold' if bold else font
        leading=leading or size*1.24
        lines=wrap(text,font,size,w)
        box=self.p.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(len(lines)*leading+5));tf=box.text_frame;tf.clear();tf.word_wrap=False
        tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if not i else tf.add_paragraph();p.text=line;p.font.name='Georgia' if font=='Georgia' else 'Arial';p.font.size=Pt(size);p.font.bold=font=='Arial-Bold';p.font.color.rgb=rgb(color);p.space_after=Pt(0);p.line_spacing=Pt(leading)
            if link:
                for run in p.runs:run.hyperlink.address=link
            self.c.setFont(font,size);self.c.setFillColor(color);self.c.drawString(x,H-y-size-i*leading,line)
        if link:self.c.linkURL(link,(x,H-y-len(lines)*leading,x+w,H-y),relative=1,thickness=0)
        return len(lines)*leading
    def image(self,file,x,y,w,h):
        with Image.open(file) as im:iw,ih=im.size
        scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale;xx=x+(w-dw)/2;yy=y+(h-dh)/2
        pic=self.p.shapes.add_picture(str(file),Pt(xx),Pt(yy),width=Pt(dw),height=Pt(dh))
        pic._element.nvPicPr.cNvPr.set('descr','Luo Heritage interface screenshot with labelled practice details: '+file.stem)
        self.c.drawImage(str(file),xx,H-yy-dh,width=dw,height=dh,mask='auto')
    def footer(self,role,page,total):
        self.rect(36,508,888,1,'#d7d9ce')
        self.text('LUO HERITAGE  /  '+role.upper(),36,516,690,8.5,MUTED)
        self.text('Contents',788,516,62,8.5,GREEN)
        self.p.shapes[-1].name='Contents link'
        self.c.linkAbsolute('Contents','s3',(786,12,848,28),thickness=0)
        self.text(f'{page:02d} / {total:02d}',870,514,58,10,MUTED,bold=True)
    def notes(self,text): self.p.notes_slide.notes_text_frame.text=text

def draw_cover(d,deck,total):
    d.rect(0,0,W,H,BROWN);d.rect(0,0,18,H,GOLD)
    d.image(LOGO,58,45,90,90)
    d.text('LUO HERITAGE  /  SELF-STUDY SLIDES',174,61,690,13,GOLD,bold=True)
    d.text('LANGUAGE • CULTURE • LIVING MEMORY',174,89,690,11,'#e7d8bf')
    title=deck['title'];size=43 if len(title)<28 else 37
    used=d.text(title,58,177,830,size,'#fffaf0',font='Georgia',leading=size*1.17)
    d.text(deck['subtitle'],60,192+used,795,21,'#f3e6ce',leading=28)
    d.rect(58,358,840,94,'#6d4930',radius=8)
    d.text(deck['goal'],78,375,795,19,'#fff7e8',leading=25)
    d.text('READ • TRY • CHECK',60,485,550,12,GOLD,bold=True)
    d.text('25 October 2026  |  Edition 1.0',640,486,285,11,'#e7d8bf')
    d.notes('Self-study guide. Screenshots show actual interface components rendered with clearly labelled sample data, not a production acceptance demonstration. No real account or cultural record was changed to make these slides.')

def draw_page(d,deck,s,page,total):
    d.text('SELF-STUDY  /  '+deck['title'].upper(),36,22,885,10,MUTED,bold=True)
    title_size=27 if len(s['title'])<50 else 24
    d.text(s['title'],36,47,888,title_size,BROWN,font='Georgia',leading=30)
    d.rect(36,94,50,4,GOLD)
    shot=SHOTS/(s['image']+'.png') if s.get('image') else None
    if shot:
        assert shot.exists(),shot
        with Image.open(shot) as im:portrait=im.height/im.width>1.05
        image_box=(604,115,310,340) if portrait else (444,116,480,326)
        text_w=530 if portrait else 367
        d.rect(image_box[0]-8,image_box[1]-8,image_box[2]+16,image_box[3]+16,'#e5e3d8',radius=7)
        d.image(shot,*image_box)
        d.text('CURRENT INTERFACE • PRACTICE DETAILS',image_box[0],image_box[1]+image_box[3]+12,image_box[2],8.2,MUTED)
    else:text_w=875
    y=119
    for n,text in enumerate(s['steps'],1):
        d.rect(36,y+1,26,26,GREEN,radius=5);d.text(str(n),44,y+3,18,15,'#ffffff',bold=True)
        height=d.text(text,76,y,text_w-40,17 if shot else 20,INK,leading=23 if shot else 27)
        y+=max(height,29)+16
    if s.get('diagram'):
        labels=s['diagram'];count=len(labels);boxw=(864-(count-1)*28)/count
        diag_y=max(y+5,322)
        assert diag_y<407,(deck['stem'],s['title'],'diagram overflow',diag_y)
        for i,label in enumerate(labels):
            x=44+i*(boxw+28);d.rect(x,diag_y,boxw,59,GREEN,radius=6)
            lines=wrap(label,'Arial-Bold',16,boxw-20);d.text(label,x+12,diag_y+(59-len(lines)*20)/2,boxw-20,16,'#ffffff',bold=True,leading=20)
            if i<count-1:d.text('→',x+boxw+5,diag_y+17,28,20,BROWN,bold=True)
        y=diag_y+73
    # A consistent result bar tells self-directed readers what to look for.
    result=s.get('result')
    if result:
        result_y=464
        maxwidth=888
        linecount=len(wrap('CHECK: '+result,'Arial-Bold',11.5,maxwidth-24))
        if result_y+linecount*15>502:result_y=459
        d.rect(36,result_y,888,linecount*15+11,'#e5eee6',radius=5)
        d.text('CHECK: '+result,48,result_y+5,864,11.5,GREEN,bold=True,leading=15)
    if s.get('tip'):
        # Tips are part of the visible slide, not hidden presenter notes.
        tipw=text_w-2 if shot else 870
        tip_size=11.2
        tip_y=y+1 if not s.get('diagram') else y+1
        tip_height=len(wrap(s['tip'],'Arial',tip_size,tipw))*14
        result_limit=458 if result else 495
        if tip_y+tip_height+6 <= result_limit:
            d.text(s['tip'],42,tip_y,tipw,tip_size,MUTED,leading=14)
        else:
            # Keep long tips accessible in speaker notes and the separate source;
            # adding them here would crowd the screenshot or result box.
            raise ValueError(('Visible tip does not fit',deck['stem'],page,s['title'],tip_y,tip_height))
    d.footer(deck['title'],page,total)
    d.notes('\n\n'.join([s['title']]+[f'{i}. {t}' for i,t in enumerate(s['steps'],1)]+(['Expected result: '+result] if result else [])+(['Important: '+s['tip']] if s.get('tip') else [])+([s['notes']] if s.get('notes') else [])+(['Screenshot: current application interface rendered with sample details. No live action is demonstrated.'] if shot else [])))

def export(deck):
    slides=deck['slides'];total=len(slides)+1
    p=Presentation();p.slide_width=Pt(W);p.slide_height=Pt(H)
    p.core_properties.title='Luo Heritage — '+deck['title'];p.core_properties.subject='Self-study training slides';p.core_properties.author='Luo Heritage';p.core_properties.keywords='self-study, role guide, illustrated, training'
    pdf=canvas.Canvas(str(OUT/(deck['stem']+'.pdf')),pagesize=(W,H));pdf.setTitle('Luo Heritage — '+deck['title']);pdf.setAuthor('Luo Heritage')
    d=SlideDraw(p,pdf);draw_cover(d,deck,total);pdf.bookmarkPage('cover');pdf.addOutlineEntry(deck['title'],'cover',0);pdf.showPage()
    for idx,s in enumerate(slides,2):
        d=SlideDraw(p,pdf);draw_page(d,deck,s,idx,total);pdf.bookmarkPage('s'+str(idx));pdf.addOutlineEntry(s['title'],'s'+str(idx),1);pdf.showPage()
    for page in p.slides:
        for shape in page.shapes:
            if shape.name=='Contents link':shape.click_action.target_slide=p.slides[2]
    pdf.save();p.save(str(OUT/(deck['stem']+'.pptx')))
    return dict(file=deck['stem'],role=deck['title'],slides=total,screenshotSlides=sum(bool(s.get('image')) for s in slides),diagrams=sum(bool(s.get('diagram')) for s in slides))

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    stats=[export(d) for d in [START]+decks]
    (OUT/'source/slide-content.json').write_text(json.dumps([START]+decks,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'source/build-summary.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
    for name in ['capture-luo-training.cjs','build_luo_training.py']:
        shutil.copy2(Path(__file__).parent/name,OUT/'source'/name)
    rows=''.join(f'<tr><td>{i["role"]}</td><td>{i["slides"]}</td><td><a href="{i["file"]}.pptx">PowerPoint</a></td><td><a href="{i["file"]}.pdf">PDF slides</a></td></tr>' for i in stats)
    (OUT/'START_HERE.html').write_text('''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Luo Heritage training slides</title><style>body{font:17px/1.6 Arial;color:#252723;background:#f7f3ea;margin:0}main{max-width:1000px;margin:auto;padding:40px 22px}h1{font:38px Georgia;color:#55351e}table{border-collapse:collapse;width:100%;background:white}td,th{text-align:left;padding:14px;border-bottom:1px solid #ddd}th{background:#55351e;color:white}a{color:#245e4d;font-weight:bold}p{max-width:800px}.table{overflow-x:auto}</style><main><h1>Luo Heritage<br>Self-study training slides</h1><p>Choose the guide for your assigned duty. Each PowerPoint has a matching PDF of the same slides. Read a step, try it in your account, then check the result.</p><p>Screenshots show the current interface with clearly labelled sample details. They do not show real records or prove that a live action was completed. Menus depend on assigned access.</p><div class="table"><table><thead><tr><th>Guide</th><th>Slides</th><th>Editable slides</th><th>Read and share</th></tr></thead><tbody>'''+rows+'</tbody></table></div><p>Edition 1.0 · 25 October 2026. Keep passwords and private research material out of support messages. Ask the person who supplied your account if you need help.</p></main></html>',encoding='utf-8')
    (OUT/'README.txt').write_text('LUO HERITAGE — SELF-STUDY SLIDE LIBRARY\n\nOpen START_HERE.html or 00_START_HERE.pdf.\nEach numbered role has an editable .pptx and a matching .pdf.\nShare only the guide(s) relevant to each person.\nThe PDFs are slide decks, not a long-form manual.\n\nScreenshots: current application components rendered with practice data, not live-account or provider acceptance evidence. No real accounts, passwords or research text were used.\nSource files and screenshot-to-component references are in source/.\nThere are eleven role guides plus a short selection deck.\n\nEdition 1.0 — 25 October 2026\n',encoding='utf-8')
    archive=OUT.parent/'Luo-Heritage-Training-Slides.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for file in OUT.iterdir():
            if file.is_file() and file.suffix in ['.pdf','.pptx','.html','.txt']:z.write(file,'Luo-Heritage-Training/'+file.name)
    print(json.dumps(dict(decks=len(stats),slides=sum(i['slides'] for i in stats),roles=len(decks),files=stats,zip=str(archive)),indent=2))

if __name__=='__main__': main()
