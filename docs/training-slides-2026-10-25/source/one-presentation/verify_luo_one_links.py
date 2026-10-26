import sys,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent/'training-python'))
from pypdf import PdfReader
from pptx import Presentation
root=Path('D:/my-business-website/deliverables/Luo-Heritage-One-Presentation')
index=json.loads((root/'section-index.json').read_text(encoding='utf8'))
pdf=PdfReader(root/'Luo-Heritage-Complete-Training.pdf')
page_ids={page.indirect_reference.idnum for page in pdf.pages}
links=0
for page in pdf.pages:
    for item in page.get('/Annots',[]):
        annotation=item.get_object()
        destination=annotation.get('/Dest')
        action=annotation.get('/A')
        assert not action or action.get('/S') not in ['/GoToR','/Launch'],'External file action'
        if destination:
            assert destination[0].idnum in page_ids,'Invalid internal page target'
            links+=1
assert links>=149,links
pres=Presentation(root/'Luo-Heritage-Complete-Training.pptx')
targets=[]
for shape in pres.slides[2].shapes:
    try:
        target=shape.click_action.target_slide
    except ValueError:
        continue
    if target is not None:
        targets.append(next(i+1 for i,slide in enumerate(pres.slides) if slide.slide_id==target.slide_id))
expected={3,4}|{section['start'] for section in index['roles']}
assert set(targets)==expected,(set(targets),expected)
assert len(targets)==37,len(targets)
report=json.loads((root/'verification.json').read_text(encoding='utf8'))
report.update(pdfInternalLinksChecked=links,powerPointContentsLinksChecked=len(targets))
(root/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf8')
print(json.dumps({'pdfLinks':links,'powerPointContentsLinks':len(targets),'roles':len(index['roles'])}))
