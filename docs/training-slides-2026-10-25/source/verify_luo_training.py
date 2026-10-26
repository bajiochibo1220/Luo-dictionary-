import sys,json,re,zipfile,hashlib
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent/'training-python'))
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image, ImageDraw
ROOT=Path('D:/my-business-website/deliverables/Luo-Heritage-Training')
DATA=json.loads((ROOT/'source/slide-content.json').read_text(encoding='utf-8'))
SUMMARY=json.loads((ROOT/'source/build-summary.json').read_text(encoding='utf-8'))
normalize=lambda s:re.sub(r'\s+',' ',s).strip()
results=[]
previews=ROOT/'source/quality-previews';previews.mkdir(exist_ok=True)
for guide,summary in zip(DATA,SUMMARY):
    stem=guide['stem'];pres=Presentation(ROOT/(stem+'.pptx'));pdf=PdfReader(ROOT/(stem+'.pdf'))
    assert len(pdf.pages)==len(pres.slides)==len(guide['slides'])+1==summary['slides'],stem
    assert len(pdf.outline)>=1,stem+' PDF bookmarks missing'
    for i,page in enumerate(pdf.pages):
        assert abs(float(page.mediabox.width)-960)<.1 and abs(float(page.mediabox.height)-540)<.1
        pp=pres.slides[i];pdftext=normalize(page.extract_text());ppttext=normalize(' '.join(shape.text for shape in pp.shapes if shape.has_text_frame))
        assert all(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=pres.slide_width+10000 and shape.top+shape.height<=pres.slide_height+10000 for shape in pp.shapes),(stem,i,'shape outside slide')
        if i:
            expected=guide['slides'][i-1]
            for text in [expected['title']]+expected['steps']+([expected['tip']] if expected.get('tip') else []):
                assert normalize(text) in pdftext,(stem,i+1,'PDF text missing',text)
                assert normalize(text) in ppttext,(stem,i+1,'PowerPoint text missing',text)
            if expected.get('image'):
                pics=[shape for shape in pp.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE]
                assert len(pics)==1,(stem,i,'Missing embedded screenshot')
                assert pics[0]._element.nvPicPr.cNvPr.get('descr'),(stem,i,'Missing screenshot description')
            contents=[shape for shape in pp.shapes if shape.name=='Contents link']
            assert len(contents)==1 and contents[0].click_action.target_slide==pres.slides[2],(stem,i,'Broken contents link')
        for font_resource in page.get('/Resources',{}).get('/Font',{}).values():
            pass
        assert not any(term in pdftext.lower() for term in ['lorem ipsum','coming soon','coming in phase','tbd']),(stem,i,'Unfinished copy')
    with zipfile.ZipFile(ROOT/(stem+'.pptx')) as z:
        assert z.testzip() is None
        assert any(n.startswith('ppt/media/') for n in z.namelist())
        for name in z.namelist():
            if name.endswith('.rels'):
                relationships=z.read(name).decode('utf-8')
                assert not re.search(r'Type="[^\"]*/image"[^>]*TargetMode="External"',relationships)
    document=pdfium.PdfDocument(str(ROOT/(stem+'.pdf')))
    # Render every page, not only a selection, to detect malformed PDF content.
    thumbs=[]
    for i in range(len(document)):
        page=document[i];im=page.render(scale=.32).to_pil().convert('RGB')
        assert im.getbbox(),(stem,i,'Empty rendered page')
        thumb=Image.new('RGB',(320,204),'#ffffff');thumb.paste(im,(6,3));ImageDraw.Draw(thumb).text((8,181),f'{stem} | slide {i+1}',fill='#252723');thumbs.append(thumb)
        if i in [0, min(6,len(document)-1),len(document)-1]:page.render(scale=1.25).to_pil().save(previews/f'{stem}-{i+1:02}.png')
    columns=4;rows=(len(thumbs)+columns-1)//columns;contact=Image.new('RGB',(columns*320,rows*204),'#dfded6')
    for i,thumb in enumerate(thumbs):contact.paste(thumb,((i%columns)*320,(i//columns)*204))
    contact.save(previews/(stem+'-contact.png'))
    results.append({'guide':stem,'slides':len(pdf.pages),'allVisibleInstructions':True,'screenshotsEmbedded':True,'contentsLinksChecked':True,'allPDFPagesRendered':True})
    print('PASS',stem,len(pdf.pages),'slides',flush=True)
archive=ROOT.parent/'Luo-Heritage-Training-Slides.zip'
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    assert len([n for n in z.namelist() if n.endswith('.pptx')])==12
    assert len([n for n in z.namelist() if n.endswith('.pdf')])==12
    for n in z.namelist():
        local=ROOT/Path(n).name
        assert hashlib.sha256(z.read(n)).hexdigest()==hashlib.sha256(local.read_bytes()).hexdigest()
report={'date':'2026-10-25','guides':len(results),'roles':11,'slides':sum(r['slides'] for r in results),'checks':results,'archiveChecked':True,'limits':'Interface screenshots use isolated sample data, not live authentication, provider delivery, publication, or real-device acceptance. PowerPoint packages were structurally checked; visual previews were rendered from matching PDFs.'}
(ROOT/'source/verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='checks'},indent=2))
